"""Emite um DARF no Sicalc com dados de quatro células de uma planilha .xlsx."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import sys
import time
from urllib.parse import urlsplit

SICALC_URL = "https://sicalc.receita.fazenda.gov.br/sicalc/principal"


@dataclass(frozen=True)
class DadosDarf:
    codigo: str
    valor: Decimal
    periodo: str  # DD/MM/AAAA: DD é o decêndio 01, 02 ou 03.
    cnpj: str


def normalizar_codigo(valor: object) -> str:
    texto = str(valor).strip()
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) != 6 or not re.fullmatch(r"[\d\s-]+", texto):
        raise ValueError("Código deve ter seis dígitos, como 3467-02 ou 346702.")
    return f"{digitos[:4]}-{digitos[4:]}"


def normalizar_valor(valor: object) -> Decimal:
    if isinstance(valor, bool) or valor is None:
        raise ValueError("Valor principal ausente ou inválido.")
    if isinstance(valor, str):
        texto = valor.strip().replace("R$", "").replace(" ", "")
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d{1,3}(\.\d{3})+", texto):
            texto = texto.replace(".", "")
    else:
        texto = str(valor)
    try:
        numero = Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f"Valor principal inválido: {valor!r}.") from exc
    if (
        not numero.is_finite()
        or numero <= 0
        or numero != numero.quantize(Decimal("0.01"))
    ):
        raise ValueError(
            "Valor principal deve ser positivo e ter até duas casas decimais."
        )
    return numero.quantize(Decimal("0.01"))


def normalizar_periodo(valor: object) -> str:
    if isinstance(valor, (date, datetime)):
        dia, mes, ano = valor.day, valor.month, valor.year
    else:
        texto = str(valor).strip()
        digitos = re.sub(r"\D", "", texto)
        if not re.fullmatch(r"[\d/\s]+", texto) or len(digitos) not in (7, 8):
            raise ValueError("Período deve ser DDMMAAAA, por exemplo 03092026.")
        digitos = digitos.zfill(8)  # Excel pode eliminar o zero inicial.
        dia, mes, ano = int(digitos[:2]), int(digitos[2:4]), int(digitos[4:])
    try:
        date(ano, mes, dia)
    except ValueError as exc:
        raise ValueError("Data do período de apuração inválida.") from exc
    return f"{dia:02d}/{mes:02d}/{ano:04d}"


def normalizar_cnpj(valor: object) -> str:
    texto = str(valor).strip()
    if not re.fullmatch(r"[\d./\s-]+", texto):
        raise ValueError("CNPJ contém caracteres inválidos.")
    digitos = re.sub(r"\D", "", texto)
    if len(digitos) != 14 or len(set(digitos)) == 1:
        raise ValueError("CNPJ deve conter 14 dígitos válidos.")

    def dv(parte: str, pesos: list[int]) -> str:
        resto = sum(int(d) * p for d, p in zip(parte, pesos)) % 11
        return str(0 if resto < 2 else 11 - resto)

    primeiro = dv(digitos[:12], [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    segundo = dv(digitos[:12] + primeiro, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
    if digitos[-2:] != primeiro + segundo:
        raise ValueError("Dígitos verificadores do CNPJ inválidos.")
    return digitos


def criar_dados(
    codigo: object, valor: object, periodo: object, cnpj: object
) -> DadosDarf:
    dados = DadosDarf(
        normalizar_codigo(codigo),
        normalizar_valor(valor),
        normalizar_periodo(periodo),
        normalizar_cnpj(cnpj),
    )
    if dados.codigo == "3467-02" and int(dados.periodo[:2]) not in (1, 2, 3):
        raise ValueError(
            "Para 3467-02, o período deve iniciar pelo decêndio 01, 02 ou 03."
        )
    return dados


def ler_planilha(caminho: Path, referencias: list[str]) -> DadosDarf:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_from_string
    except ImportError as exc:
        raise RuntimeError(
            "Instale as dependências: python -m pip install -r requirements.txt"
        ) from exc
    if not caminho.is_file() or caminho.suffix.lower() != ".xlsx":
        raise ValueError("Informe uma planilha .xlsx existente em --planilha.")
    try:
        arquivo = load_workbook(caminho, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Não foi possível abrir a planilha: {caminho}.") from exc
    try:
        valores = []
        for referencia in referencias:
            if "!" not in referencia:
                raise ValueError(
                    f"Referência {referencia!r} deve ter formato aba!célula."
                )
            aba, celula = referencia.rsplit("!", 1)
            aba, celula = aba.strip("'\" "), celula.strip().upper()
            try:
                coordinate_from_string(celula)
            except Exception as exc:
                raise ValueError(f"Célula inválida em {referencia!r}.") from exc
            if aba.isdigit():
                indice = int(aba)
                if not 1 <= indice <= len(arquivo.worksheets):
                    raise ValueError(f"Aba {indice} não existe na planilha.")
                folha = arquivo.worksheets[indice - 1]
            elif aba in arquivo.sheetnames:
                folha = arquivo[aba]
            else:
                raise ValueError(f"Aba {aba!r} não existe na planilha.")
            valor = folha[celula].value
            if valor is None:
                raise ValueError(
                    f"{referencia} está vazia (ou tem fórmula sem resultado salvo)."
                )
            valores.append(valor)
        return criar_dados(*valores)
    finally:
        arquivo.close()


def moeda_sicalc(valor: Decimal) -> str:
    return f"{valor:.2f}".replace(".", ",")


def mostrar_chrome() -> None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="chrome", headless=False)
        try:
            page = browser.new_page()
            page.goto(SICALC_URL, wait_until="domcontentloaded")
            print("Chrome aberto no Sicalc. Pressione Enter para fechar.")
            try:
                input()
            except EOFError:
                page.wait_for_timeout(30000)
        finally:
            browser.close()


def selecionar_codigo(page, codigo: str) -> None:
    page.locator("#codReceitaPrincipal").fill(codigo)
    opcoes = page.locator(
        ".ui-autocomplete:visible li, [role='listbox']:visible [role='option']"
    )
    try:
        opcoes.first.wait_for(state="visible", timeout=15000)
    except Exception as exc:
        raise RuntimeError("A lista de códigos do Sicalc não apareceu.") from exc
    padrao = re.compile(
        r"^\s*"
        + re.escape(codigo[:4])
        + r"\s*[-/]?\s*"
        + re.escape(codigo[-2:])
        + r"\b"
    )
    for indice in range(opcoes.count()):
        opcao = opcoes.nth(indice)
        if padrao.search(opcao.inner_text()):
            opcao.click()
            return
    raise RuntimeError(f"O Sicalc não ofereceu o código exato {codigo}.")


def campo_valor_principal(page):
    # O ID do campo não foi fornecido; o rótulo visível é o seletor principal.
    por_rotulo = page.get_by_label(re.compile(r"valor\s+(do\s+)?principal", re.I))
    if por_rotulo.count() == 1:
        return por_rotulo
    por_id = page.locator(
        "#valorPrincipal, #valorPrincipalImposto, input[name='valorPrincipal']"
    )
    if por_id.count() == 1:
        return por_id
    raise RuntimeError(
        "Não foi possível identificar unicamente o campo Valor Principal."
    )


def resumo_apos_continuar(page, cnpj: str) -> str:
    conteudo = ""
    for seletor in ("#fldError", "[role='alert']", "[role='main']", "body"):
        try:
            campo = page.locator(seletor).first
            if campo.count():
                if seletor == "#fldError":
                    conteudo = (campo.text_content(timeout=2000) or "").strip()
                elif campo.is_visible():
                    conteudo = campo.inner_text(timeout=2000).strip()
            if conteudo:
                break
        except Exception:
            continue
    if not conteudo:
        conteudo = "(a página não apresentou uma mensagem legível)"
    conteudo = " ".join(conteudo.split())[:1000]
    conteudo = conteudo.replace(cnpj, "***")
    conteudo = re.sub(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b", "***", conteudo)
    return f"URL: {page.url.replace(cnpj, '***')}; conteúdo: {conteudo}"


def finalizar_darf(page, context, dados: DadosDarf, destino: Path) -> Path:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError

    selecionar_codigo(page, dados.codigo)
    periodo = page.locator("#dataPA")
    periodo.fill(dados.periodo)
    periodo.press("Tab")  # Dispara o blur que calcula o vencimento.
    try:
        page.wait_for_function(
            "() => !!document.querySelector('#dataVencimentoImposto')?.value",
            timeout=15000,
        )
    except PlaywrightTimeoutError as exc:
        raise RuntimeError("O Sicalc não calculou a data de vencimento.") from exc
    principal = campo_valor_principal(page)
    principal.fill(moeda_sicalc(dados.valor))
    principal.press("Tab")
    if normalizar_valor(principal.input_value()) != dados.valor:
        raise RuntimeError("Valor principal exibido diverge da planilha.")
    page.locator("#btnCalcular").click()
    selecao = page.locator(
        "input[type='checkbox'][data-darf='true'][data-pagamento='true']"
    )
    selecao.first.wait_for(state="visible", timeout=30000)
    if selecao.count() != 1:
        raise RuntimeError("O cálculo não retornou exatamente um DARF selecionável.")
    selecao.check()
    downloads = []
    respostas_pdf = []
    page.on("download", downloads.append)
    context.on("page", lambda nova: nova.on("download", downloads.append))

    def capturar_pdf(resposta):
        if "application/pdf" not in resposta.headers.get("content-type", "").lower():
            return
        try:
            corpo = resposta.body()
            if corpo.startswith(b"%PDF"):
                respostas_pdf.append(corpo)
        except Exception:
            pass

    context.on("response", capturar_pdf)
    page.locator("#btnDarf").click()
    limite = time.monotonic() + 45
    while not downloads and not respostas_pdf and time.monotonic() < limite:
        page.wait_for_timeout(250)
    if downloads:
        downloads[0].save_as(destino)
    elif respostas_pdf:
        destino.write_bytes(respostas_pdf[0])
    else:
        raise RuntimeError("O Sicalc não entregou um PDF em 45 segundos.")
    if destino.read_bytes()[:4] != b"%PDF":
        destino.unlink(missing_ok=True)
        raise RuntimeError("O arquivo recebido não é um PDF válido.")
    return destino


def emitir_darf(dados: DadosDarf, saida: Path, timeout_captcha: int) -> Path:
    try:
        from playwright.sync_api import Error as PlaywrightError
        from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
        from playwright.sync_api import sync_playwright
    except ImportError as exc:
        raise RuntimeError(
            "Instale as dependências: python -m pip install -r requirements.txt"
        ) from exc
    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / (
        f"darf_{dados.codigo}_{dados.periodo.replace('/', '')}_{datetime.now():%Y%m%d_%H%M%S_%f}.pdf"
    )
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="chrome", headless=False)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        try:
            page.goto(SICALC_URL, wait_until="domcontentloaded")
            page.locator("#linkRapido").click()
            page.locator("#optionPJ").check()
            page.locator("#cnpjFormatado").fill(dados.cnpj)
            continuar = page.locator("input[type='submit'][value='Continuar']")
            try:
                continuar.wait_for(state="visible", timeout=timeout_captcha)
                print(
                    "Marque 'Sou humano' no Chrome. Aguardando pelo menos 10 segundos..."
                )
                page.wait_for_timeout(10_000)
                print("Aguardando o botão Continuar ficar habilitado...")
                page.wait_for_function(
                    "() => { const botao = document.querySelector(\"input[type='submit'][value='Continuar']\"); return botao && !botao.disabled && window.sicalc?.getHCaptchaResponse()?.trim(); }",
                    timeout=timeout_captcha,
                )
            except PlaywrightTimeoutError as exc:
                raise RuntimeError(
                    "O botão Continuar não foi habilitado com uma resposta do hCaptcha."
                ) from exc
            dialogos = []
            requisicoes = []
            respostas_captcha = []
            falhas_captcha = []
            respostas_formulario = []
            erros_javascript = []
            envio_captcha = {}

            def limpar_mensagem_resposta(conteudo):
                texto = re.sub(r"<[^>]*>", " ", conteudo)
                texto = " ".join(texto.split())[:500]
                texto = texto.replace(dados.cnpj, "***")
                texto = re.sub(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b", "***", texto)
                return re.sub(r"[A-Za-z0-9_-]{32,}", "[omitido]", texto)

            def registrar_dialogo(dialogo):
                dialogos.append(f"{dialogo.type}: {dialogo.message}")
                dialogo.dismiss()

            def registrar_requisicao(requisicao):
                if "sicalc.receita.fazenda.gov.br" in requisicao.url:
                    caminho = urlsplit(requisicao.url).path
                    requisicoes.append(f"{requisicao.method} {caminho}")
                    if requisicao.method == "POST" and caminho.endswith("/api/captcha"):
                        try:
                            corpo = requisicao.post_data_json
                            token = corpo.get("h-recaptcha-response", "")
                            envio_captcha.update(
                                {
                                    "token_presente": bool(token),
                                    "cnpj_confere": re.sub(
                                        r"\D", "", corpo.get("cnpjFormatado", "")
                                    )
                                    == dados.cnpj,
                                }
                            )
                        except (AttributeError, TypeError, ValueError):
                            envio_captcha.update(
                                {"dados": "não foi possível inspecionar"}
                            )

            def registrar_resposta(resposta):
                if resposta.request.method != "POST":
                    return
                caminho = urlsplit(resposta.url).path
                if caminho.endswith("/api/captcha"):
                    try:
                        mensagem = limpar_mensagem_resposta(resposta.text())
                    except PlaywrightError:
                        mensagem = "(não foi possível ler a resposta)"
                    respostas_captcha.append((resposta.status, mensagem))
                    print(f"Validação do CAPTCHA pelo Sicalc: HTTP {resposta.status}.")
                elif caminho.endswith("/rapido/contribuinte"):
                    respostas_formulario.append(resposta.status)

            def registrar_falha(requisicao):
                if urlsplit(requisicao.url).path.endswith("/api/captcha"):
                    falhas_captcha.append(str(requisicao.failure or "falha de rede"))

            page.on("dialog", registrar_dialogo)
            page.on("request", registrar_requisicao)
            page.on("response", registrar_resposta)
            page.on("requestfailed", registrar_falha)
            page.on("pageerror", lambda erro: erros_javascript.append(str(erro)))
            page.evaluate("""() => {
                    const formulario = document.querySelector('#contribuinte');
                    window.__sicalcDiagnostico = { envios: 0, invalidos: [] };
                    formulario.addEventListener('submit', () => {
                        window.__sicalcDiagnostico.envios++;
                    }, true);
                    formulario.addEventListener('invalid', (evento) => {
                        window.__sicalcDiagnostico.invalidos.push(evento.target.id);
                    }, true);
                }""")
            print("Botão Continuar habilitado e CAPTCHA preenchido. Clicando.")
            continuar.click()
            if dialogos:
                raise RuntimeError(
                    "O Sicalc exibiu um diálogo ao clicar em Continuar: "
                    + " | ".join(dialogos)[:500]
                )
            limite_codigo = time.monotonic() + 30
            while time.monotonic() < limite_codigo:
                pagina_codigo = next(
                    (
                        pagina
                        for pagina in context.pages
                        if pagina.locator("#codReceitaPrincipal").is_visible()
                    ),
                    None,
                )
                if pagina_codigo is not None:
                    page = pagina_codigo
                    break
                if respostas_captcha and respostas_captcha[-1][0] != 200:
                    break
                page.wait_for_timeout(250)
            else:
                pagina_codigo = None
            if pagina_codigo is None:
                captcha = (
                    f"HTTP {respostas_captcha[-1][0]}; resposta: {respostas_captcha[-1][1] or '(vazia)'}"
                    if respostas_captcha
                    else (
                        f"falha de rede: {falhas_captcha[-1]}"
                        if falhas_captcha
                        else "nenhuma resposta observada"
                    )
                )
                formulario_post = (
                    f"HTTP {respostas_formulario[-1]}"
                    if respostas_formulario
                    else "não enviado"
                )
                eventos = ", ".join(requisicoes[-8:]) or "nenhuma requisição ao Sicalc"
                javascript = (
                    " | ".join(erros_javascript[-2:])[:400] or "nenhum erro capturado"
                )
                try:
                    formulario = page.evaluate(
                        "() => window.__sicalcDiagnostico || null"
                    )
                except PlaywrightError:
                    formulario = None
                raise RuntimeError(
                    "Após clicar em Continuar, o campo do código não apareceu. "
                    f"Validação em /api/captcha: {captcha}; envio: {envio_captcha}; "
                    f"POST do formulário: {formulario_post}. "
                    f"Formulário: {formulario}. Requisições: {eventos}. "
                    f"JavaScript: {javascript}. "
                    + resumo_apos_continuar(page, dados.cnpj)
                )
            return finalizar_darf(page, context, dados, destino)
        except PlaywrightError as exc:
            raise RuntimeError(f"Falha na interação com o Sicalc: {exc}") from exc
        finally:
            context.close()
            browser.close()


def emitir_darf_apos_entrada_manual(
    dados: DadosDarf, saida: Path, downloads: Path, inspecionar: bool = False
) -> Path | None:
    from interface_windows import continuar_no_chrome_comum

    saida.mkdir(parents=True, exist_ok=True)
    destino = saida / (
        f"darf_{dados.codigo}_{dados.periodo.replace('/', '')}_{datetime.now():%Y%m%d_%H%M%S_%f}.pdf"
    )
    print(f"Abra no Chrome comum: {SICALC_URL}")
    print(
        "Depois: Preenchimento rápido → Pessoa Jurídica → CNPJ → "
        "Sou humano → Continuar."
    )
    print("Quando aparecer o campo do código, volte ao terminal e pressione Enter.")
    try:
        input()
    except EOFError as exc:
        raise RuntimeError("O modo manual requer um terminal interativo.") from exc
    print("Clique na aba do Sicalc no Chrome. A automação começará em 5 segundos...")
    time.sleep(5)
    try:
        return continuar_no_chrome_comum(
            codigo=dados.codigo,
            valor=dados.valor,
            periodo=dados.periodo,
            cnpj=dados.cnpj,
            downloads=downloads,
            destino=destino,
            inspecionar=inspecionar,
        )
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError(f"Falha no controle da interface do Chrome: {exc}") from exc


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Lê quatro células de um .xlsx e emite um DARF."
    )
    parser.add_argument(
        "referencias",
        nargs="*",
        metavar="ABA!CÉLULA",
        help="código, valor, período e CNPJ, nessa ordem",
    )
    parser.add_argument("--planilha", type=Path, help="arquivo .xlsx")
    parser.add_argument(
        "--exemplo", action="store_true", help="usa os valores fornecidos na conversa"
    )
    parser.add_argument(
        "--abrir-chrome",
        action="store_true",
        help="abre somente o Sicalc no Chrome, sem preencher dados",
    )
    parser.add_argument(
        "--entrada-manual",
        action="store_true",
        help="faça CNPJ/CAPTCHA no Chrome comum; depois controle a interface do Windows",
    )
    parser.add_argument(
        "--inspecionar-interface",
        action="store_true",
        help="no modo manual, lista controles acessíveis sem emitir DARF",
    )
    parser.add_argument(
        "--downloads",
        type=Path,
        default=Path.home() / "Downloads",
        help="pasta de downloads configurada no Chrome",
    )
    parser.add_argument(
        "--validar", action="store_true", help="só valida dados; não acessa o Sicalc"
    )
    parser.add_argument("--saida", type=Path, default=Path.cwd(), help="pasta do PDF")
    parser.add_argument(
        "--timeout-captcha",
        type=int,
        default=300,
        help="segundos para verificação humana",
    )
    args = parser.parse_args()
    try:
        if args.abrir_chrome:
            if (
                args.exemplo
                or args.planilha
                or args.referencias
                or args.validar
                or args.entrada_manual
            ):
                parser.error("--abrir-chrome deve ser usado sozinho.")
            mostrar_chrome()
            return 0
        if args.exemplo:
            if args.planilha or args.referencias:
                parser.error(
                    "--exemplo não pode ser combinado com planilha ou referências."
                )
            dados = criar_dados("3467-02", "20", "03092026", "55642790000109")
        else:
            if not args.planilha or len(args.referencias) != 4:
                parser.error(
                    "Informe --planilha e quatro referências: código, valor, período e CNPJ."
                )
            dados = ler_planilha(args.planilha, args.referencias)
        print(
            f"Código: {dados.codigo}; principal: R$ {moeda_sicalc(dados.valor)}; "
            f"período: {dados.periodo}; CNPJ: ***{dados.cnpj[-4:]}"
        )
        if args.validar:
            print("Dados válidos. Nenhum acesso ao Sicalc foi feito.")
            return 0
        if args.timeout_captcha <= 0:
            raise ValueError("--timeout-captcha deve ser maior que zero.")
        if args.inspecionar_interface and not args.entrada_manual:
            parser.error("--inspecionar-interface requer --entrada-manual.")
        if args.entrada_manual:
            destino = emitir_darf_apos_entrada_manual(
                dados, args.saida, args.downloads, args.inspecionar_interface
            )
        else:
            destino = emitir_darf(dados, args.saida, args.timeout_captcha * 1000)
        if destino is not None:
            print(f"PDF salvo em: {destino}")
        return 0
    except (ValueError, RuntimeError, OSError) as exc:
        print(f"Erro: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
