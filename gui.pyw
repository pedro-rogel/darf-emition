"""Interface gráfica para emitir um DARF com entrada manual no Chrome comum."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from queue import Empty, Queue
import os
import shutil
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from darf import DadosDarf, SICALC_URL, ler_planilha, moeda_sicalc
from interface_windows import ativar_janela_sicalc, continuar_no_chrome_comum

BASE_DIR = Path(__file__).resolve().parent


def localizar_chrome() -> Path:
    candidatos = [
        shutil.which("chrome.exe"),
        *(
            str(Path(base) / "Google" / "Chrome" / "Application" / "chrome.exe")
            for base in (
                os.environ.get("PROGRAMFILES"),
                os.environ.get("PROGRAMFILES(X86)"),
                os.environ.get("LOCALAPPDATA"),
            )
            if base
        ),
    ]
    for candidato in candidatos:
        if candidato and Path(candidato).is_file():
            return Path(candidato)
    raise RuntimeError("Google Chrome não encontrado neste computador.")


def nome_do_pdf(dados: DadosDarf) -> str:
    return (
        f"darf_{dados.codigo}_{dados.periodo.replace('/', '')}_"
        f"{datetime.now():%Y%m%d_%H%M%S_%f}.pdf"
    )


class Aplicativo(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Emissão de DARF — Sicalc")
        largura = min(920, max(700, self.winfo_screenwidth() - 80))
        altura = min(650, max(500, self.winfo_screenheight() - 100))
        self.geometry(f"{largura}x{altura}")
        self.minsize(700, 500)
        self.configure(bg="#f4f7fb")
        self._configurar_estilo()

        self.planilha = tk.StringVar()
        self.saida = tk.StringVar(value=str(BASE_DIR))
        self.downloads = tk.StringVar(value=str(Path.home() / "Downloads"))
        self.abas = [tk.StringVar() for _ in range(4)]
        self.celulas = [tk.StringVar() for _ in range(4)]
        self.status = tk.StringVar(value="Selecione a planilha para começar.")
        self.resumo = tk.StringVar(value="Nenhum dado conferido ainda.")
        self.pdf = tk.StringVar()
        self.dados: DadosDarf | None = None
        self.referencias: tuple[str, ...] | None = None
        self.em_execucao = False
        self.eventos: Queue[tuple[str, object]] = Queue()

        self._montar()
        for variavel in (
            self.planilha,
            self.saida,
            self.downloads,
            *self.abas,
            *self.celulas,
        ):
            variavel.trace_add("write", self._invalidar)
        self.after(100, self._processar_eventos)
        self.protocol("WM_DELETE_WINDOW", self._fechar)

    def _configurar_estilo(self) -> None:
        estilo = ttk.Style(self)
        if "vista" in estilo.theme_names():
            estilo.theme_use("vista")
        estilo.configure("TLabel", font=("Segoe UI", 10), background="#ffffff")
        estilo.configure("TEntry", font=("Segoe UI", 10))
        estilo.configure("TCombobox", font=("Segoe UI", 10))
        estilo.configure("TButton", font=("Segoe UI", 10), padding=(12, 7))
        estilo.configure("TLabelframe", background="#ffffff", padding=14)
        estilo.configure(
            "TLabelframe.Label", font=("Segoe UI Semibold", 11), background="#ffffff"
        )

    def _montar(self) -> None:
        cabecalho = tk.Frame(self, bg="#17365d", padx=24, pady=18)
        cabecalho.pack(fill="x")
        tk.Label(
            cabecalho,
            text="Emissão de DARF",
            bg="#17365d",
            fg="white",
            font=("Segoe UI Semibold", 20),
        ).pack(anchor="w")
        tk.Label(
            cabecalho,
            text="Escolha as células de um .xlsx e conclua a identificação no Chrome.",
            bg="#17365d",
            fg="#dce8f4",
            font=("Segoe UI", 10),
        ).pack(anchor="w", pady=(4, 0))

        area_rolagem = tk.Frame(self, bg="#f4f7fb")
        area_rolagem.pack(fill="both", expand=True)
        self.tela = tk.Canvas(
            area_rolagem, bg="#f4f7fb", borderwidth=0, highlightthickness=0
        )
        barra = ttk.Scrollbar(area_rolagem, orient="vertical", command=self.tela.yview)
        self.tela.configure(yscrollcommand=barra.set)
        self.tela.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        corpo = tk.Frame(self.tela, bg="#f4f7fb", padx=22, pady=18)
        self.corpo_na_tela = self.tela.create_window((0, 0), window=corpo, anchor="nw")
        corpo.bind(
            "<Configure>",
            lambda _evento: self.tela.configure(scrollregion=self.tela.bbox("all")),
        )
        self.tela.bind("<Configure>", self._ajustar_largura)
        self.tela.bind_all("<MouseWheel>", self._rolar)

        planilha_frame = ttk.LabelFrame(corpo, text="1  Planilha .xlsx")
        planilha_frame.pack(fill="x", pady=(0, 12))
        planilha_frame.columnconfigure(0, weight=1)
        ttk.Label(
            planilha_frame, text="Arquivo .xlsx do Excel ou baixado do Google Sheets"
        ).grid(row=0, column=0, sticky="w")
        ttk.Entry(planilha_frame, textvariable=self.planilha).grid(
            row=1, column=0, sticky="ew", pady=(5, 0), padx=(0, 8)
        )
        ttk.Button(
            planilha_frame, text="Procurar...", command=self._procurar_planilha
        ).grid(row=1, column=1, pady=(5, 0), padx=(0, 8))
        ttk.Button(
            planilha_frame, text="Carregar abas", command=self._carregar_abas
        ).grid(row=1, column=2, pady=(5, 0))

        celulas_frame = ttk.LabelFrame(corpo, text="2  Onde estão os dados?")
        celulas_frame.pack(fill="x", pady=(0, 12))
        celulas_frame.columnconfigure(1, weight=1)
        ttk.Label(celulas_frame, text="Informação").grid(
            row=0, column=0, sticky="w", padx=(0, 18)
        )
        ttk.Label(celulas_frame, text="Aba da planilha").grid(
            row=0, column=1, sticky="w"
        )
        ttk.Label(celulas_frame, text="Célula").grid(
            row=0, column=2, sticky="w", padx=(12, 0)
        )
        nomes = ("Código do imposto", "Valor principal", "Período de apuração", "CNPJ")
        self.seletores_abas = []
        for indice, nome in enumerate(nomes, start=1):
            ttk.Label(celulas_frame, text=nome).grid(
                row=indice, column=0, sticky="w", pady=6, padx=(0, 18)
            )
            seletor = ttk.Combobox(
                celulas_frame,
                textvariable=self.abas[indice - 1],
                state="readonly",
                width=38,
            )
            seletor.grid(row=indice, column=1, sticky="ew", pady=6)
            self.seletores_abas.append(seletor)
            ttk.Entry(
                celulas_frame, textvariable=self.celulas[indice - 1], width=12
            ).grid(row=indice, column=2, sticky="w", pady=6, padx=(12, 0))
        ttk.Label(
            celulas_frame,
            text="Exemplo: terceira aba + A8 = código em 3!A8.",
            foreground="#576b80",
        ).grid(row=5, column=0, columnspan=3, sticky="w", pady=(10, 0))

        destino_frame = ttk.LabelFrame(corpo, text="3  Onde salvar o PDF?")
        destino_frame.pack(fill="x", pady=(0, 12))
        destino_frame.columnconfigure(0, weight=1)
        ttk.Label(destino_frame, text="Pasta para o arquivo final").grid(
            row=0, column=0, sticky="w"
        )
        ttk.Entry(destino_frame, textvariable=self.saida).grid(
            row=1, column=0, sticky="ew", pady=(5, 0), padx=(0, 8)
        )
        ttk.Button(
            destino_frame,
            text="Escolher pasta...",
            command=lambda: self._procurar_pasta(self.saida),
        ).grid(row=1, column=1, pady=(5, 0))
        ttk.Label(
            destino_frame,
            text="Pasta de downloads do Chrome (altere somente se for diferente da padrão)",
            foreground="#576b80",
        ).grid(row=2, column=0, sticky="w", pady=(12, 0))
        ttk.Entry(destino_frame, textvariable=self.downloads).grid(
            row=3, column=0, sticky="ew", pady=(5, 0), padx=(0, 8)
        )
        ttk.Button(
            destino_frame,
            text="Escolher pasta...",
            command=lambda: self._procurar_pasta(self.downloads),
        ).grid(row=3, column=1, pady=(5, 0))

        acoes = ttk.Frame(corpo)
        acoes.pack(fill="x", pady=(0, 12))
        ttk.Button(acoes, text="Conferir dados", command=self._conferir).pack(
            side="left"
        )
        ttk.Button(acoes, text="Abrir Sicalc no Chrome", command=self._abrir_site).pack(
            side="left", padx=(9, 0)
        )
        self.botao_cnpj = ttk.Button(
            acoes, text="Copiar CNPJ", command=self._copiar_cnpj, state="disabled"
        )
        self.botao_cnpj.pack(side="left", padx=(9, 0))
        self.botao_continuar = ttk.Button(
            acoes,
            text="Já estou na tela do código — emitir DARF",
            command=self._continuar,
            state="disabled",
        )
        self.botao_continuar.pack(side="right")

        resumo_frame = ttk.LabelFrame(corpo, text="Conferência e andamento")
        resumo_frame.pack(fill="x")
        ttk.Label(resumo_frame, textvariable=self.resumo, wraplength=820).pack(
            anchor="w"
        )
        ttk.Label(
            resumo_frame, textvariable=self.status, wraplength=820, foreground="#17365d"
        ).pack(anchor="w", pady=(10, 0))
        self.botao_pdf = ttk.Button(
            resumo_frame,
            text="Abrir PDF salvo",
            command=self._abrir_pdf,
            state="disabled",
        )
        self.botao_pdf.pack(anchor="w", pady=(12, 0))

        rodape = tk.Label(
            self,
            text="Etapa manual: Pessoa Jurídica → CNPJ → Sou humano → Continuar. "
            "Depois volte aqui e clique no botão de emitir.",
            bg="#e9f1f8",
            fg="#17365d",
            font=("Segoe UI", 10),
            padx=22,
            pady=12,
            anchor="w",
            wraplength=850,
        )
        rodape.pack(fill="x", side="bottom")

    def _ajustar_largura(self, evento) -> None:
        largura = min(evento.width, 1100)
        self.tela.itemconfigure(self.corpo_na_tela, width=largura)
        self.tela.coords(self.corpo_na_tela, max(0, (evento.width - largura) // 2), 0)

    def _rolar(self, evento) -> None:
        if self.winfo_viewable():
            self.tela.yview_scroll(-int(evento.delta / 120), "units")

    def _invalidar(self, *_args) -> None:
        self.dados = None
        self.referencias = None
        self.resumo.set("Dados alterados. Clique em “Conferir dados” novamente.")
        if not self.em_execucao:
            self.botao_continuar.configure(state="disabled")
            self.botao_cnpj.configure(state="disabled")

    def _procurar_planilha(self) -> None:
        caminho = filedialog.askopenfilename(
            parent=self,
            title="Escolha a planilha",
            filetypes=[("Planilhas Excel", "*.xlsx")],
        )
        if caminho:
            self.planilha.set(caminho)
            self._carregar_abas()

    def _carregar_abas(self) -> None:
        try:
            from openpyxl import load_workbook

            caminho = Path(self.planilha.get().strip().strip('"')).expanduser()
            if self.planilha.get().strip().startswith(("https://", "http://")):
                raise ValueError(
                    "No Google Sheets, use Arquivo > Fazer download > Microsoft Excel (.xlsx) "
                    "e escolha o arquivo baixado. Um link não pode ser usado como caminho local."
                )
            if not caminho.is_file() or caminho.suffix.lower() != ".xlsx":
                raise ValueError("Escolha uma planilha .xlsx existente.")
            arquivo = load_workbook(caminho, read_only=True, data_only=True)
            try:
                nomes = arquivo.sheetnames
            finally:
                arquivo.close()
            for seletor in self.seletores_abas:
                seletor.configure(values=nomes)
            padroes = (
                nomes[2] if len(nomes) >= 3 else "",
                nomes[1] if len(nomes) >= 2 else "",
                nomes[0] if nomes else "",
                nomes[0] if nomes else "",
            )
            for variavel, nome in zip(self.abas, padroes):
                variavel.set(nome)
            for variavel, celula in zip(self.celulas[:2], ("A8", "A9")):
                if not variavel.get().strip():
                    variavel.set(celula)
            self.status.set(
                f"Planilha carregada: {len(nomes)} aba(s). Indique as quatro células."
            )
        except Exception as exc:
            self.status.set("Não foi possível carregar a planilha.")
            messagebox.showerror("Planilha", str(exc), parent=self)

    def _procurar_pasta(self, variavel: tk.StringVar) -> None:
        inicial = Path(variavel.get()).expanduser()
        caminho = filedialog.askdirectory(
            parent=self,
            title="Escolha uma pasta",
            initialdir=str(inicial if inicial.is_dir() else BASE_DIR),
        )
        if caminho:
            variavel.set(caminho)

    def _configuracao(self) -> tuple[DadosDarf, Path, Path, tuple[str, ...]]:
        caminho = Path(self.planilha.get().strip().strip('"')).expanduser()
        if not caminho.is_file():
            raise ValueError("Escolha a planilha .xlsx antes de continuar.")
        referencias = []
        for indice, (aba, celula) in enumerate(zip(self.abas, self.celulas), start=1):
            nome_aba, endereco = aba.get().strip(), celula.get().strip().upper()
            if not nome_aba or not endereco:
                raise ValueError(f"Informe a aba e a célula da informação {indice}.")
            referencias.append(f"{nome_aba}!{endereco}")
        saida = Path(self.saida.get().strip().strip('"')).expanduser()
        downloads = Path(self.downloads.get().strip().strip('"')).expanduser()
        if not saida.is_dir():
            raise ValueError("Escolha uma pasta existente para salvar o PDF final.")
        if not downloads.is_dir():
            raise ValueError(
                "A pasta de downloads do Chrome não existe. Escolha a pasta correta."
            )
        dados = ler_planilha(caminho, referencias)
        return dados, saida, downloads, tuple(referencias)

    def _conferir(self) -> None:
        try:
            dados, _saida, _downloads, referencias = self._configuracao()
        except Exception as exc:
            self.status.set("Corrija as informações e tente novamente.")
            messagebox.showerror("Conferência dos dados", str(exc), parent=self)
            return
        self.dados = dados
        self.referencias = referencias
        self.resumo.set(
            f"Código {dados.codigo}  •  Valor R$ {moeda_sicalc(dados.valor)}  •  "
            f"Período {dados.periodo}  •  CNPJ {dados.cnpj}"
        )
        self.status.set(
            "Dados conferidos. No Chrome, conclua Pessoa Jurídica, CNPJ, “Sou humano” "
            "e Continuar. Depois clique em “Já estou na tela do código”."
        )
        self.botao_continuar.configure(state="normal")
        self.botao_cnpj.configure(state="normal")
        self.after_idle(lambda: self.tela.yview_moveto(1))

    def _copiar_cnpj(self) -> None:
        if self.dados is not None:
            self.clipboard_clear()
            self.clipboard_append(self.dados.cnpj)
            self.status.set("CNPJ copiado. Cole no campo do Sicalc.")

    def _abrir_site(self) -> None:
        try:
            subprocess.Popen(
                [str(localizar_chrome()), SICALC_URL],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
            self.status.set(
                "Sicalc aberto no Chrome. Faça a etapa manual até a tela do código."
            )
        except Exception as exc:
            messagebox.showerror("Google Chrome", str(exc), parent=self)

    def _continuar(self) -> None:
        if self.em_execucao or self.dados is None:
            return
        try:
            dados, saida, downloads, referencias = self._configuracao()
            if referencias != self.referencias or dados != self.dados:
                raise ValueError(
                    "Os dados mudaram. Clique em “Conferir dados” novamente."
                )
        except Exception as exc:
            messagebox.showerror("Configuração", str(exc), parent=self)
            return
        self.em_execucao = True
        self.botao_continuar.configure(state="disabled")
        self.botao_pdf.configure(state="disabled")
        self.pdf.set("")
        self.status.set("Localizando a guia do Sicalc no Chrome...")
        self.withdraw()
        threading.Thread(
            target=self._executar, args=(dados, saida, downloads), daemon=True
        ).start()

    def _executar(self, dados: DadosDarf, saida: Path, downloads: Path) -> None:
        pythoncom = None
        com_inicializado = False
        try:
            import pythoncom as modulo_com

            pythoncom = modulo_com
            pythoncom.CoInitialize()
            com_inicializado = True
            saida.mkdir(parents=True, exist_ok=True)
            ativar_janela_sicalc()
            destino = saida / nome_do_pdf(dados)
            resultado = continuar_no_chrome_comum(
                codigo=dados.codigo,
                valor=dados.valor,
                periodo=dados.periodo,
                cnpj=dados.cnpj,
                downloads=downloads,
                destino=destino,
                progresso=lambda texto: self.eventos.put(("progresso", texto)),
            )
            if resultado is None:
                raise RuntimeError("O Sicalc não entregou um PDF.")
            self.eventos.put(("concluido", resultado))
        except Exception as exc:
            self.eventos.put(("erro", str(exc)))
        finally:
            if com_inicializado:
                pythoncom.CoUninitialize()

    def _processar_eventos(self) -> None:
        try:
            while True:
                tipo, valor = self.eventos.get_nowait()
                if tipo == "progresso":
                    self.status.set(str(valor))
                elif tipo == "concluido":
                    self.em_execucao = False
                    self.deiconify()
                    self.lift()
                    self.tela.yview_moveto(1)
                    self.pdf.set(str(valor))
                    self.botao_pdf.configure(state="normal")
                    self.status.set(f"PDF salvo em: {valor}")
                    messagebox.showinfo(
                        "DARF salvo", f"PDF salvo em:\n{valor}", parent=self
                    )
                elif tipo == "erro":
                    self.em_execucao = False
                    self.deiconify()
                    self.lift()
                    self.tela.yview_moveto(1)
                    self.botao_continuar.configure(state="normal")
                    self.status.set(
                        "A emissão não foi concluída. Confira a mensagem de erro."
                    )
                    messagebox.showerror(
                        "Não foi possível concluir", str(valor), parent=self
                    )
        except Empty:
            pass
        self.after(100, self._processar_eventos)

    def _abrir_pdf(self) -> None:
        caminho = Path(self.pdf.get())
        if caminho.is_file():
            os.startfile(caminho)

    def _fechar(self) -> None:
        if self.em_execucao:
            messagebox.showwarning(
                "Emissão em andamento", "Aguarde a emissão terminar.", parent=self
            )
            return
        self.destroy()


def main() -> None:
    if sys.platform != "win32":
        raise RuntimeError("A interface gráfica usa o controle de janelas do Windows.")
    Aplicativo().mainloop()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        raiz = tk.Tk()
        raiz.withdraw()
        messagebox.showerror("Emissão de DARF", str(exc), parent=raiz)
        raiz.destroy()
